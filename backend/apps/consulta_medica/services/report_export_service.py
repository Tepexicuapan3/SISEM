import io

import openpyxl

# Encabezados en el mismo orden que las columnas devueltas por
# daily_report_usecase._consultation_to_report_row -- si se agrega una
# columna ahi, agregarla tambien aqui.
_DAILY_REPORT_COLUMNS = (
    ("date", "Fecha"),
    ("time", "Hora"),
    ("folio", "Folio"),
    ("noExp", "No. Expediente"),
    ("pkNum", "Familiar (0=titular)"),
    ("patientName", "Paciente"),
    ("consultorio", "Consultorio"),
    ("doctorName", "Medico"),
    ("primaryDiagnosis", "Diagnostico"),
    ("cieCode", "Clave CIE"),
    ("cieDescription", "Descripcion CIE"),
    ("serviceType", "Tipo de Servicio"),
)


def build_daily_report_workbook(items: list[dict]) -> bytes:
    """
    Genera un .xlsx real (openpyxl) -- a diferencia del legado
    (dataexcel.jsp/nexcel.jsp), que solo ponia un content-type de Excel
    sobre una tabla HTML. Ver docs/architecture/legacy-reports-inventory.md.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Informe Diario Consulta"

    for col_idx, (_, header) in enumerate(_DAILY_REPORT_COLUMNS, start=1):
        ws.cell(row=1, column=col_idx, value=header)

    for row_idx, item in enumerate(items, start=2):
        for col_idx, (field, _) in enumerate(_DAILY_REPORT_COLUMNS, start=1):
            value = item.get(field)
            # openpyxl no acepta date/time con tzinfo naive-mixto de forma
            # consistente entre columnas -- se normaliza a texto para evitar
            # sorpresas de formato en Excel.
            if value is not None and not isinstance(value, (str, int, float, bool)):
                value = str(value)
            ws.cell(row=row_idx, column=col_idx, value=value)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
