import io

import openpyxl

# Encabezados en el mismo orden que
# MedicalLeaveRepository.to_report_row -- si se agrega una columna ahi,
# agregarla tambien aqui.
_MEDICAL_LEAVE_REPORT_COLUMNS = (
    ("startDate", "Fecha Inicio"),
    ("endDate", "Fecha Fin"),
    ("days", "Dias"),
    ("folio", "Folio"),
    ("noExp", "No. Expediente"),
    ("pkNum", "Familiar (0=titular)"),
    ("patientName", "Paciente"),
    ("doctorName", "Medico"),
    ("leaveTypeName", "Tipo de Licencia"),
    ("isSubsequent", "Subsecuente"),
)


def build_medical_leave_report_workbook(items: list[dict]) -> bytes:
    """
    Genera un .xlsx real (openpyxl) -- mismo criterio que
    consulta_medica/services/report_export_service.py (informe diario) y
    pases/services/report_export_service.py.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Informe de Incapacidades"

    for col_idx, (_, header) in enumerate(_MEDICAL_LEAVE_REPORT_COLUMNS, start=1):
        ws.cell(row=1, column=col_idx, value=header)

    for row_idx, item in enumerate(items, start=2):
        for col_idx, (field, _) in enumerate(_MEDICAL_LEAVE_REPORT_COLUMNS, start=1):
            value = item.get(field)
            if isinstance(value, bool):
                value = "Sí" if value else "No"
            elif value is not None and not isinstance(value, (str, int, float)):
                value = str(value)
            ws.cell(row=row_idx, column=col_idx, value=value)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
