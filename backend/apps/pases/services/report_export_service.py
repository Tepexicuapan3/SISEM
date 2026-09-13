import io

import openpyxl

# Encabezados en el mismo orden que ReferralRepository.to_report_row -- si se
# agrega una columna ahi, agregarla tambien aqui.
_REFERRAL_REPORT_COLUMNS = (
    ("date", "Fecha"),
    ("folio", "Folio"),
    ("referralType", "Tipo de Pase"),
    ("noExp", "No. Expediente"),
    ("pkNum", "Familiar (0=titular)"),
    ("patientName", "Paciente"),
    ("doctorName", "Medico"),
    ("destinationCenterName", "Destino"),
    ("specialtyName", "Especialidad"),
    ("visitType", "Tipo de Cita"),
    ("status", "Estatus"),
)


def build_referral_report_workbook(items: list[dict]) -> bytes:
    """
    Genera un .xlsx real (openpyxl) -- mismo criterio que
    consulta_medica/services/report_export_service.py, en vez de replicar el
    truco de dataexcel.jsp del legado (HTML con content-type de Excel).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Informe de Pases"

    for col_idx, (_, header) in enumerate(_REFERRAL_REPORT_COLUMNS, start=1):
        ws.cell(row=1, column=col_idx, value=header)

    for row_idx, item in enumerate(items, start=2):
        for col_idx, (field, _) in enumerate(_REFERRAL_REPORT_COLUMNS, start=1):
            value = item.get(field)
            if value is not None and not isinstance(value, (str, int, float, bool)):
                value = str(value)
            ws.cell(row=row_idx, column=col_idx, value=value)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
