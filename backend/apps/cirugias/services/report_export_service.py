import io

import openpyxl

# Encabezados en el mismo orden que SurgeryRepository.to_report_row -- si se
# agrega una columna ahi, agregarla tambien aqui.
_SURGERY_REPORT_COLUMNS = (
    ("date", "Fecha"),
    ("time", "Hora"),
    ("folio", "Folio"),
    ("noExp", "No. Expediente"),
    ("pkNum", "Familiar (0=titular)"),
    ("surgeonName", "Cirujano"),
    ("surgeryTypeName", "Tipo de Cirugia"),
    ("classificationName", "Clasificacion"),
    ("status", "Estatus"),
    ("performedStatus", "Realizacion"),
)


def build_surgery_report_workbook(items: list[dict]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Informe de Cirugias"

    for col_idx, (_, header) in enumerate(_SURGERY_REPORT_COLUMNS, start=1):
        ws.cell(row=1, column=col_idx, value=header)

    for row_idx, item in enumerate(items, start=2):
        for col_idx, (field, _) in enumerate(_SURGERY_REPORT_COLUMNS, start=1):
            value = item.get(field)
            if value is not None and not isinstance(value, (str, int, float, bool)):
                value = str(value)
            ws.cell(row=row_idx, column=col_idx, value=value)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
