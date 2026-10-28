"""Excel de la auditoria de solicitudes (openpyxl), mismo criterio que los informes de pases."""
import io

import openpyxl

# Mismo orden que las filas de AuditoriaUseCase: si se agrega un campo alla, agregarlo aqui.
COLUMNAS = (
    ("folio", "Folio"),
    ("folioOrigen", "Folio de origen"),
    ("tipoNombre", "Tipo"),
    ("centroNombre", "Centro"),
    ("clues", "CLUES"),
    ("estatus", "Estatus"),
    ("modalidad", "Modalidad"),
    ("aprobaciones", "Aprobaciones vigentes"),
    ("aprobacionesRequeridas", "Aprobaciones requeridas"),
    ("totalItems", "Conceptos"),
    ("solicitanteNombre", "Solicitó"),
    ("fechaCreacion", "Creada"),
    ("fechaLimite", "Vence"),
    ("fechaEscalamiento", "Escalada"),
    ("fechaResolucion", "Resuelta"),
    ("resolvioNombre", "Resolvió"),
    ("motivoCierre", "Motivo de cierre"),
)


def _valor(valor):
    if valor is None or isinstance(valor, (str, int, float, bool)):
        return valor
    if hasattr(valor, "tzinfo") and getattr(valor, "tzinfo", None) is not None:
        from django.utils import timezone

        return timezone.localtime(valor).strftime("%Y-%m-%d %H:%M")
    return str(valor)


def libro_auditoria(filas: list[dict]) -> bytes:
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Auditoría de autorizaciones"
    for columna, (_, encabezado) in enumerate(COLUMNAS, start=1):
        hoja.cell(row=1, column=columna, value=encabezado)
    for renglon, fila in enumerate(filas, start=2):
        for columna, (campo, _) in enumerate(COLUMNAS, start=1):
            hoja.cell(row=renglon, column=columna, value=_valor(fila.get(campo)))
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
