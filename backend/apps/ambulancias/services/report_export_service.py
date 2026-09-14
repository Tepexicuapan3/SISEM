import io

import openpyxl

# Encabezados en el mismo orden que AmbulanceRequestRepository.to_report_row
# -- si se agrega una columna ahi, agregarla tambien aqui.
_AMBULANCE_REPORT_COLUMNS = (
    ("date", "Fecha de Registro"),
    ("folio", "Folio"),
    ("noExp", "No. Expediente"),
    ("pkNum", "Familiar (0=titular)"),
    ("requestingClinicName", "Clinica Solicitante"),
    ("reasonName", "Motivo"),
    ("destinationName", "Destino"),
    ("transferDate", "Fecha de Traslado"),
    ("status", "Estatus"),
    ("authorizationStatus", "Autorizacion"),
    ("serviceNumber", "No. Servicio"),
)


def build_ambulance_report_workbook(items: list[dict]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Informe de Ambulancias"

    for col_idx, (_, header) in enumerate(_AMBULANCE_REPORT_COLUMNS, start=1):
        ws.cell(row=1, column=col_idx, value=header)

    for row_idx, item in enumerate(items, start=2):
        for col_idx, (field, _) in enumerate(_AMBULANCE_REPORT_COLUMNS, start=1):
            value = item.get(field)
            if value is not None and not isinstance(value, (str, int, float, bool)):
                value = str(value)
            ws.cell(row=row_idx, column=col_idx, value=value)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
