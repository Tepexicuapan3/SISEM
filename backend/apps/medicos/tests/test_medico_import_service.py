"""Cobertura de `medico_import_service.parse_and_validate` (R6, medicos-id-no-medico):
No. Médico, usuario, nombres, cédula, Escuela / Especialidad (warnings,
D7), Estatus A/B y duplicados por nombre. Usa la BD de pruebas (SQLite)."""

import io
from datetime import date

import openpyxl
from django.contrib.auth.hashers import make_password
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase

from apps.authentication.models import DetUsuario, SyUsuario
from apps.catalogos.models import CatTipoPersonal, Escuelas, Especialidades
from apps.medicos.models import CatMedico
from apps.medicos.services.medico_import_service import (
    HEADERS,
    ImportFileError,
    build_template,
    parse_and_validate,
)


def row(**kw):
    """Fila del Excel en el orden de HEADERS; las claves son abreviaturas."""
    keys = [
        "no", "usuario", "nom", "pat", "mat", "ced", "esc", "sexo", "fnac",
        "esp", "serv", "estatus",
    ]
    unknown = set(kw) - set(keys)
    assert not unknown, unknown
    return tuple(kw.get(k, "") for k in keys)


def xlsx(rows, filename="import.xlsx", headers=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(headers or HEADERS))
    for r in rows:
        ws.append(list(r))
    buffer = io.BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile(filename, buffer.getvalue())


def make_usuario(username, tipo_personal="Médico"):
    """Usuario con DetUsuario; `tipo_personal=None` lo deja sin tipo (no es médico)."""
    user = SyUsuario.objects.create(
        usuario=username,
        correo=f"{username}@example.com",
        clave_hash=make_password("x"),
        est_activo=True,
    )
    tipo = (
        CatTipoPersonal.objects.get_or_create(name=tipo_personal, defaults={"is_active": True})[0]
        if tipo_personal
        else None
    )
    DetUsuario.objects.create(id_usuario=user, nombre=username, id_tipo_personal=tipo)
    return user


class MedicoImportParseTests(TestCase):
    def _parse(self, *rows):
        return parse_and_validate(xlsx(rows))

    # ── Archivo ─────────────────────────────────────────────────────────────

    def test_headers_antiguos_se_rechazan(self):
        with self.assertRaises(ImportFileError) as ctx:
            parse_and_validate(
                xlsx([], headers=["ID del Médico", "Usuario (Login del sistema)"])
            )
        self.assertEqual(ctx.exception.code, "IMPORT_HEADERS_MISMATCH")

    def test_plantilla_tiene_headers_y_tres_filas_de_ejemplo(self):
        wb = openpyxl.load_workbook(io.BytesIO(build_template()))
        ws = wb.active
        self.assertEqual([c.value for c in ws[1]], HEADERS)
        self.assertEqual(ws.max_row, 4)
        # Una con usuario, una sin usuario con cédula, una sin No. Médico.
        self.assertTrue(ws.cell(2, 2).value)
        self.assertFalse(ws.cell(3, 2).value)
        self.assertTrue(ws.cell(3, 6).value)
        self.assertFalse(ws.cell(4, 1).value)

    # ── No. Médico (S1.4, R1) ───────────────────────────────────────────────

    def test_no_medico_valido_se_parsea_a_entero(self):
        result = self._parse(row(no="2726", pat="LOPEZ", nom="JUAN", ced="123"))
        self.assertEqual(result["total_errores"], 0)
        self.assertEqual(result["rows"][0]["data"]["noMedico"], 2726)

    def test_no_medico_con_sufijo_decimal_de_excel_se_acepta(self):
        result = self._parse(row(no="35.0", pat="LOPEZ", nom="JUAN", ced="123"))
        self.assertEqual(result["rows"][0]["data"]["noMedico"], 35)

    def test_no_medico_invalido_es_error(self):
        for raw in ("ABC", "0", "-5", "1.5", "12A"):
            with self.subTest(raw=raw):
                result = self._parse(row(no=raw, pat="LOPEZ", nom="JUAN", ced="1"))
                self.assertEqual(result["total_errores"], 1)
                self.assertIn("no es un entero positivo", result["rows"][0]["errors"][0])

    def test_no_medico_duplicado_en_archivo_es_error(self):
        result = self._parse(
            row(no="10", pat="A", nom="UNO", ced="1"),
            row(no="10", pat="B", nom="DOS", ced="1"),
        )
        self.assertEqual(result["total_errores"], 2)
        for r in result["rows"]:
            self.assertIn("No. Médico 10 duplicado en el archivo.", r["errors"])

    def test_no_medico_existente_en_bd_es_error(self):
        CatMedico.objects.create(id=77)
        result = self._parse(row(no="77", pat="A", nom="UNO", ced="1"))
        self.assertEqual(result["total_errores"], 1)
        self.assertIn("Ya existe un médico con el No. Médico 77.", result["rows"][0]["errors"])

    # ── Usuario (S6.1, S6.4) ────────────────────────────────────────────────

    def test_con_usuario_ignora_nombres_y_avisa(self):
        make_usuario("dr_uno")
        result = self._parse(row(usuario="dr_uno", pat="X", nom="Y", ced="999"))
        r = result["rows"][0]
        self.assertEqual(r["errors"], [])
        self.assertEqual(len(r["warnings"]), 1)
        self.assertIsNone(r["data"]["paterno"])
        self.assertIsNone(r["data"]["nombre"])
        self.assertEqual(r["data"]["cedulaProfesional"], "999")

    def test_usuario_inexistente_es_error(self):
        result = self._parse(row(usuario="no_existe"))
        self.assertEqual(result["total_errores"], 1)
        self.assertIn("No existe un usuario", result["rows"][0]["errors"][0])

    def test_usuario_no_medico_es_error(self):
        # S6.4: misma regla que el alta manual (tipo_personal = Médico).
        make_usuario("admin_no_medico", tipo_personal="Administrativo")
        make_usuario("sin_tipo", tipo_personal=None)
        result = self._parse(row(usuario="admin_no_medico"), row(usuario="sin_tipo"))
        self.assertEqual(result["total_errores"], 2)
        for r in result["rows"]:
            self.assertIn("no tiene tipo de personal Médico", r["errors"][0])

    def test_usuario_sin_detalle_es_error(self):
        SyUsuario.objects.create(
            usuario="sin_detalle", correo="sin_detalle@example.com",
            clave_hash=make_password("x"), est_activo=True,
        )
        result = self._parse(row(usuario="sin_detalle"))
        self.assertIn("no tiene tipo de personal Médico", result["rows"][0]["errors"][0])

    def test_usuario_que_ya_es_medico_es_error(self):
        user = make_usuario("ya_medico")
        CatMedico.objects.create(id_usuario=user)
        result = self._parse(row(usuario="ya_medico"))
        self.assertIn("Este usuario ya tiene un perfil de médico.", result["rows"][0]["errors"])

    def test_usuario_repetido_en_archivo_es_error(self):
        make_usuario("repetido")
        result = self._parse(row(usuario="repetido"), row(usuario="repetido"))
        self.assertEqual(result["total_errores"], 2)

    # ── Sin usuario (S6.2, S6.3) ────────────────────────────────────────────

    def test_sin_usuario_exige_paterno_y_nombre(self):
        result = self._parse(row(ced="1"))
        errors = result["rows"][0]["errors"]
        self.assertIn("Nombre(s) es obligatorio cuando no se captura Usuario.", errors)
        self.assertIn("Apellido Paterno es obligatorio cuando no se captura Usuario.", errors)

    def test_sin_usuario_con_datos_y_cedula_sin_avisos(self):
        result = self._parse(row(pat="López", mat="Ramírez", nom="Juan", ced="1234567"))
        r = result["rows"][0]
        self.assertEqual((r["errors"], r["warnings"]), ([], []))
        self.assertEqual(r["data"]["paterno"], "López")
        self.assertEqual(r["data"]["cedulaProfesional"], "1234567")

    def test_sin_usuario_sin_cedula_es_warning_no_error(self):
        result = self._parse(row(pat="López", nom="Juan"))
        r = result["rows"][0]
        self.assertEqual(r["errors"], [])
        self.assertEqual(len(r["warnings"]), 1)
        self.assertEqual(result["total_errores"], 0)
        self.assertEqual(result["total_warnings"], 1)

    def test_cedula_de_mas_de_30_caracteres_es_error(self):
        result = self._parse(row(pat="A", nom="B", ced="9" * 31))
        self.assertEqual(result["total_errores"], 1)

    def test_duplicado_por_nombre_en_archivo_solo_sin_no_medico(self):
        result = self._parse(
            row(pat="López", mat="R", nom="Juan", ced="1"),
            row(pat="LOPEZ", mat="R", nom="JUAN", ced="1"),
        )
        self.assertEqual(result["total_errores"], 2)

    def test_mismo_nombre_con_no_medico_no_se_considera_duplicado(self):
        result = self._parse(
            row(no="1", pat="López", nom="Juan", ced="1"),
            row(no="2", pat="López", nom="Juan", ced="1"),
        )
        self.assertEqual(result["total_errores"], 0)

    def test_nombre_existente_en_bd_sin_no_medico_es_error(self):
        CatMedico.objects.create(paterno="Pérez", materno="G", nombre="Ana")
        result = self._parse(row(pat="PEREZ", mat="g", nom="ana", ced="1"))
        self.assertEqual(result["total_errores"], 1)
        self.assertIn("Ya existe un médico sin usuario", result["rows"][0]["errors"][0])

    # ── Especialidad (S6.6, D7) ─────────────────────────────────────────────

    def test_plantilla_ya_no_trae_clave_clinica(self):
        self.assertNotIn("Clave Clínica (legado)", HEADERS)

    def test_especialidad_por_nombre_es_error(self):
        # FK a catálogo: solo se acepta el ID.
        Especialidades.objects.create(name="Cirugía")
        r = self._parse(row(pat="A", nom="B", ced="1", esp="Cirugía"))["rows"][0]
        self.assertIsNone(r["data"]["especialidadId"])
        self.assertIn("ID numérico", r["errors"][0])

    def test_especialidad_por_id_del_catalogo(self):
        esp = Especialidades.objects.create(name="Cardiología")
        result = self._parse(row(pat="A", nom="B", ced="1", esp=str(esp.id)))
        r = result["rows"][0]
        self.assertEqual(r["data"]["especialidadId"], esp.id)
        self.assertEqual(r["warnings"], [])

    def test_especialidad_por_id_float_de_excel(self):
        esp = Especialidades.objects.create(name="Cardiología")
        result = self._parse(row(pat="A", nom="B", ced="1", esp=f"{esp.id}.0"))
        self.assertEqual(result["rows"][0]["data"]["especialidadId"], esp.id)

    def test_especialidad_id_inexistente_es_warning(self):
        result = self._parse(row(pat="A", nom="B", ced="1", esp="999999"))
        r = result["rows"][0]
        self.assertEqual(r["errors"], [])
        self.assertIsNone(r["data"]["especialidadId"])
        self.assertIn("ID 999999", r["warnings"][0])

    # ── Escuela (solo sin usuario) ──────────────────────────────────────────

    def test_escuela_por_id_del_catalogo(self):
        esc = Escuelas.objects.create(code="UNAM", name="Universidad Nacional Autónoma de México")
        result = self._parse(
            row(no="1", pat="A", nom="B", ced="1", esc=str(esc.id)),
            row(no="2", pat="C", nom="D", ced="1", esc=f"{esc.id}.0"),
        )
        self.assertEqual([r["data"]["escuelaId"] for r in result["rows"]], [esc.id] * 2)
        self.assertEqual(result["total_warnings"], 0)

    def test_escuela_id_inexistente_es_warning(self):
        r = self._parse(row(pat="A", nom="B", ced="1", esc="999999"))["rows"][0]
        self.assertEqual(r["errors"], [])
        self.assertIsNone(r["data"]["escuelaId"])
        self.assertIn("ID 999999", r["warnings"][0])

    def test_escuela_por_nombre_o_siglas_es_error(self):
        Escuelas.objects.create(code="UNAM", name="Universidad Nacional")
        result = self._parse(
            row(no="1", pat="A", nom="B", ced="1", esc="UNAM"),
            row(no="2", pat="C", nom="D", ced="1", esc="Universidad Nacional"),
        )
        self.assertEqual(result["total_errores"], 2)
        for r in result["rows"]:
            self.assertIsNone(r["data"]["escuelaId"])
            self.assertIn("ID numérico", r["errors"][0])

    def test_escuela_con_usuario_se_ignora_con_warning(self):
        make_usuario("con_escuela")
        esc = Escuelas.objects.create(code="IPN", name="Instituto Politécnico Nacional")
        r = self._parse(row(usuario="con_escuela", esc=str(esc.id)))["rows"][0]
        self.assertEqual(r["errors"], [])
        self.assertIsNone(r["data"]["escuelaId"])
        self.assertIn("Escuela del Excel se ignora", r["warnings"][0])

    # ── Sexo / Fecha de Nacimiento (solo sin usuario) ───────────────────────

    def test_sexo_acepta_codigo_o_nombre(self):
        result = self._parse(
            row(no="1", pat="A", nom="B", ced="1", sexo="M"),
            row(no="2", pat="C", nom="D", ced="1", sexo="femenino"),
            row(no="3", pat="E", nom="F", ced="1", sexo="Otro"),
            row(no="4", pat="G", nom="H", ced="1"),
        )
        self.assertEqual([r["data"]["sexo"] for r in result["rows"]], ["M", "F", "O", None])
        self.assertEqual(result["total_errores"], 0)

    def test_sexo_invalido_es_error(self):
        r = self._parse(row(pat="A", nom="B", ced="1", sexo="X"))["rows"][0]
        self.assertIsNone(r["data"]["sexo"])
        self.assertIn("Sexo debe ser", r["errors"][0])

    def test_fecha_nac_acepta_texto_y_celda_de_fecha(self):
        result = self._parse(
            row(no="1", pat="A", nom="B", ced="1", fnac="15/03/1975"),
            row(no="2", pat="C", nom="D", ced="1", fnac="1975-03-15 00:00:00"),
        )
        self.assertEqual(
            [r["data"]["fechaNac"] for r in result["rows"]], [date(1975, 3, 15)] * 2
        )

    def test_fecha_nac_invalida_o_futura_es_error(self):
        result = self._parse(
            row(no="1", pat="A", nom="B", ced="1", fnac="31/02/1980"),
            row(no="2", pat="C", nom="D", ced="1", fnac="01/01/2999"),
        )
        self.assertEqual(result["total_errores"], 2)
        self.assertIn("formato inválido", result["rows"][0]["errors"][0])
        self.assertIn("no puede ser futura", result["rows"][1]["errors"][0])
        for r in result["rows"]:
            self.assertIsNone(r["data"]["fechaNac"])

    def test_sexo_y_fecha_con_usuario_se_ignoran_con_warning(self):
        make_usuario("con_sexo")
        r = self._parse(row(usuario="con_sexo", sexo="X", fnac="basura"))["rows"][0]
        self.assertEqual(r["errors"], [])
        self.assertIsNone(r["data"]["sexo"])
        self.assertIsNone(r["data"]["fechaNac"])
        self.assertIn("Sexo y Fecha de Nacimiento del Excel se ignoran", r["warnings"][0])

    # ── Estatus / Tipo (S6.7)───────────────────────────────────────────────

    def test_estatus_a_y_b_se_mapean(self):
        result = self._parse(
            row(no="1", pat="A", nom="B", ced="1", estatus="A"),
            row(no="2", pat="A", nom="B", ced="1", estatus="b"),
            row(no="3", pat="A", nom="B", ced="1", estatus="Suspendido"),
            row(no="4", pat="A", nom="B", ced="1"),
        )
        estatus = [r["data"]["estatusMedico"] for r in result["rows"]]
        self.assertEqual(estatus, ["ACTIVO", "BAJA", "SUSPENDIDO", "ACTIVO"])
        self.assertEqual(result["total_errores"], 0)

    def test_estatus_invalido_es_error(self):
        result = self._parse(row(pat="A", nom="B", ced="1", estatus="Z"))
        self.assertEqual(result["total_errores"], 1)

    def test_tipo_medico_y_observaciones_toman_default(self):
        # La plantilla ya no trae Tipo de Médico ni Observaciones.
        result = self._parse(row(pat="A", nom="B", ced="1"))
        data = result["rows"][0]["data"]
        self.assertEqual(data["tipoMedico"], "CLINICA")
        self.assertIsNone(data["observaciones"])
