"""
Servicio de import masivo de médicos por Excel.

Calca el patrón de `apps.administracion.services.user_import_service`
(openpyxl para escribir la plantilla, pandas para leer/validar). NO persiste
nada -- ni preview ni confirm llaman a este módulo para escribir; ver
`apps.medicos.uses_case.import_medicos`, que re-corre `parse_and_validate`
sobre el MISMO archivo en ambos pasos (mecanismo stateless, sin token/cache).

Cambio `medicos-id-no-medico`: "No. Médico" (opcional) ES `CatMedico.id`.
Errores (bloquean la fila) = formato/identidad; warnings (NO bloquean) =
lookups de catálogo sin match (Clave Clínica, Especialidad) y cédula faltante
sin usuario (decisión D7).
"""

import io
import re
import unicodedata
from typing import Optional

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.authentication.models import SyUsuario
from apps.catalogos.models import CatCentroAtencion, Especialidades
from apps.medicos.models import CatMedico, ESTATUS_MEDICO, TIPO_MEDICO
from apps.medicos.repositories.medico_repository import MedicoRepository

HEADERS = [
    "No. Médico",
    "Usuario (Login del sistema)",
    "Apellido Paterno",
    "Apellido Materno",
    "Nombre(s)",
    "Cédula Profesional",
    "Clave Clínica (legado)",
    "Especialidad",
    "Servicio",
    "Tipo de Médico",
    "Estatus del Médico",
    "Observaciones",
]

MAX_ROWS = 5000

DEFAULT_TIPO_MEDICO = "CLINICA"
DEFAULT_ESTATUS_MEDICO = "ACTIVO"

# Longitudes de los campos libres de CatMedico -- se truncan defensivamente
# antes de crear para no dejar un IntegrityError sin manejar a mitad de
# `_create_all` por una celda de Excel demasiado larga. La cédula NO se
# trunca (es un identificador): si excede el máximo es error de fila.
_SERVICIO_MAX_LEN = 100
_NOMBRE_PARTE_MAX_LEN = 80
_CEDULA_MAX_LEN = 30
_MAX_NO_MEDICO = 2**63 - 1

# Alias del legado para Estatus: A=Activo, B=Baja (además de labels/códigos).
_ESTATUS_ALIASES = {"a": "ACTIVO", "b": "BAJA"}

_NO_MEDICO_RE = re.compile(r"^\d+(\.0+)?$")

_SAMPLE_ROWS = (
    # Con usuario: nombre/cédula se toman del usuario (se dejan vacíos).
    ("2726", "jperez", "", "", "", "", "12", "Cardiología", "Cardiología", "Clínica", "A", ""),
    # Sin usuario: nombre y cédula viven en el propio médico.
    ("2727", "", "LÓPEZ", "RAMÍREZ", "JUAN", "1234567", "12", "Cirugía General", "", "Hospital", "A", "Médico sin usuario del sistema"),
    # Sin No. Médico: toma el siguiente de la secuencia.
    ("", "mgomez", "", "", "", "", "", "", "", "Clínica", "A", "Alta nativa de SIRES"),
)


class ImportFileError(Exception):
    """Error de nivel-archivo (headers, extensión, límite de filas): 400, ninguna fila se evalúa."""

    def __init__(self, code: str, message: str, details: Optional[dict] = None):
        self.code = code
        self.message = message
        self.details = details or {}
        super().__init__(message)


def build_template() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Médicos"

    BRAND = "D94300"
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor=BRAND)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    for row_idx, sample_row in enumerate(_SAMPLE_ROWS, start=2):
        for col_idx, value in enumerate(sample_row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    for col_idx, header in enumerate(HEADERS, start=1):
        max_len = max(len(header), *(len(str(row[col_idx - 1])) for row in _SAMPLE_ROWS))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 4, 12), 40)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _clean_text(value) -> str:
    text = "" if value is None else str(value)
    if text.strip().lower() == "nan":
        text = ""
    return re.sub(r"\s+", " ", text.strip())


def _strip_accents(text: str) -> str:
    normalized = unicodedata.normalize("NFKD", text)
    return "".join(c for c in normalized if not unicodedata.combining(c))


def _truncate(text: str, max_length: int) -> str:
    return text[:max_length] if text else text


def _nombre_key(value) -> str:
    """Clave de comparación de nombres: sin acentos, mayúsculas, espacios colapsados."""
    return _strip_accents(_clean_text(value)).upper()


def _persona_key(paterno, materno, nombre) -> str:
    """Llave de comparación de personas sin usuario (sin acentos, mayúsculas)."""
    return _nombre_key(" ".join(p for p in (paterno, materno, nombre) if p))


def _build_choice_lookup(choices) -> dict:
    lookup = {}
    for code, label in choices:
        lookup[_strip_accents(code).lower()] = code
        lookup[_strip_accents(label).lower()] = code
    return lookup


_TIPO_MEDICO_LOOKUP = _build_choice_lookup(TIPO_MEDICO)
_ESTATUS_MEDICO_LOOKUP = {**_build_choice_lookup(ESTATUS_MEDICO), **_ESTATUS_ALIASES}


def _resolve_choice(raw_text: str, lookup: dict, default: str) -> tuple[str, bool]:
    """
    Resuelve un choice de `CatMedico` (Tipo de Médico / Estatus) por label
    o código, insensible a acentos/mayúsculas.

    Devuelve `(valor_resuelto, es_valido)`:
    - Vacío -> `(default, True)`. Comportamiento intencional y documentado
      en la plantilla/dialog: "vacío se interpreta como <default>".
    - No vacío y matchea un choice -> `(codigo, True)`.
    - No vacío y NO matchea ningún choice (typo/valor inventado) ->
      `(default, False)`. El caller (`parse_and_validate`) es responsable
      de agregar el error a la fila cuando `es_valido` es False.
    """
    if not raw_text:
        return default, True
    resolved = lookup.get(_strip_accents(raw_text).lower())
    if resolved is None:
        return default, False
    return resolved, True


def _parse_no_medico(raw: str) -> tuple[Optional[int], Optional[str]]:
    """Devuelve `(valor, error)`. Vacío -> `(None, None)` (toma la secuencia)."""
    if not raw:
        return None, None
    if not _NO_MEDICO_RE.match(raw):
        return None, f"No. Médico '{raw}' no es un entero positivo."
    value = int(raw.split(".")[0])
    if value <= 0 or value > _MAX_NO_MEDICO:
        return None, f"No. Médico '{raw}' no es un entero positivo."
    return value, None


def parse_and_validate(file) -> dict:
    """
    Única función de validación, reusada tal cual por preview y por
    confirm -- el confirm NUNCA confía en el resultado que mandó el
    cliente en preview, siempre re-corre esto contra el archivo re-subido
    (mecanismo stateless, ver `import_medicos.py`).

    Devuelve:
        {
            "total_records": int,
            "total_errores": int,     # filas con al menos un error
            "total_warnings": int,    # filas con al menos un warning
            "rows": [
                {
                    "row": int,             # numero de fila en el Excel (1-based, incluye header)
                    "data": {...},          # datos parseados en camelCase
                    "errors": [str, ...],   # bloquean la fila
                    "warnings": [str, ...], # informativos, NO bloquean
                },
                ...
            ],
        }
    """
    df = _read_dataframe(file)
    total_records = len(df)

    raw_rows = []
    for row_idx in range(total_records):
        line = df.iloc[row_idx]
        raw_rows.append(
            {
                "noMedicoRaw": _clean_text(line["No. Médico"]),
                "usuario": _clean_text(line["Usuario (Login del sistema)"]),
                "paternoRaw": _clean_text(line["Apellido Paterno"]),
                "maternoRaw": _clean_text(line["Apellido Materno"]),
                "nombreRaw": _clean_text(line["Nombre(s)"]),
                "cedulaRaw": _clean_text(line["Cédula Profesional"]),
                "claveClinicaRaw": _clean_text(line["Clave Clínica (legado)"]),
                "especialidadRaw": _clean_text(line["Especialidad"]),
                "tipoMedicoRaw": _clean_text(line["Tipo de Médico"]),
                "servicioRaw": _clean_text(line["Servicio"]),
                "estatusRaw": _clean_text(line["Estatus del Médico"]),
                "observacionesRaw": _clean_text(line["Observaciones"]),
            }
        )

    # --- No. Médico: parseo, duplicados en archivo y existentes en BD -------
    parsed_numbers = [_parse_no_medico(r["noMedicoRaw"]) for r in raw_rows]
    number_seen: dict = {}
    for value, _err in parsed_numbers:
        if value is not None:
            number_seen[value] = number_seen.get(value, 0) + 1
    existing_numbers = set(
        CatMedico.objects.filter(id__in=list(number_seen)).values_list("id", flat=True)
    )

    # --- Usuarios ------------------------------------------------------------
    usernames = [r["usuario"] for r in raw_rows if r["usuario"]]
    usuarios_by_username = {
        u.usuario: u
        for u in SyUsuario.objects.select_related("detalle__id_tipo_personal").filter(
            usuario__in=usernames
        )
    }
    existing_medico_usuario_ids = set(
        CatMedico.objects.filter(
            id_usuario_id__in=[u.id_usuario for u in usuarios_by_username.values()]
        ).values_list("id_usuario_id", flat=True)
    )
    usuario_seen: dict = {}
    for name in usernames:
        usuario_seen[name] = usuario_seen.get(name, 0) + 1

    # --- Duplicados por nombre (sin usuario y SIN No. Médico) ----------------
    # Con No. Médico la llave natural es el número; sin él, el nombre es la
    # única llave: evita que re-subir el archivo duplique las filas.
    persona_keys = []
    for r in raw_rows:
        if r["usuario"] or r["noMedicoRaw"]:
            persona_keys.append(None)
        else:
            persona_keys.append(
                _persona_key(r["paternoRaw"], r["maternoRaw"], r["nombreRaw"]) or None
            )
    persona_seen: dict = {}
    for key in persona_keys:
        if key:
            persona_seen[key] = persona_seen.get(key, 0) + 1
    existing_persona_keys = (
        {
            _persona_key(paterno or "", materno or "", nombre or "")
            for paterno, materno, nombre in CatMedico.objects.filter(
                id_usuario__isnull=True
            ).values_list("paterno", "materno", "nombre")
        }
        if any(persona_keys)
        else set()
    )

    # --- Catálogos (warning si no hay match o es ambiguo) --------------------
    claves = {int(r["claveClinicaRaw"]) for r in raw_rows if r["claveClinicaRaw"].isdigit()}
    centros_by_clave: dict = {}
    if claves:
        for centro_id, clave in CatCentroAtencion.objects.filter(
            legacy_cd_clinica__in=claves, is_active=True
        ).values_list("id", "legacy_cd_clinica"):
            centros_by_clave.setdefault(clave, []).append(centro_id)

    especialidades_by_key: dict = {}
    if any(r["especialidadRaw"] for r in raw_rows):
        for esp_id, esp_name in Especialidades.objects.filter(is_active=True).values_list(
            "id", "name"
        ):
            especialidades_by_key.setdefault(_nombre_key(esp_name), []).append(esp_id)

    result_rows = []
    total_errores = 0
    total_warnings = 0

    for row_idx, parsed in enumerate(raw_rows):
        errors = []
        warnings = []

        # No. Médico
        no_medico, number_error = parsed_numbers[row_idx]
        if number_error:
            errors.append(number_error)
        elif no_medico is not None:
            if number_seen.get(no_medico, 0) > 1:
                errors.append(f"No. Médico {no_medico} duplicado en el archivo.")
            elif no_medico in existing_numbers:
                errors.append(f"Ya existe un médico con el No. Médico {no_medico}.")

        # Usuario / identidad
        usuario_obj = None
        if not parsed["usuario"]:
            if not parsed["nombreRaw"]:
                errors.append("Nombre(s) es obligatorio cuando no se captura Usuario.")
            if not parsed["paternoRaw"]:
                errors.append("Apellido Paterno es obligatorio cuando no se captura Usuario.")
            persona_key = persona_keys[row_idx]
            if persona_key:
                if persona_seen.get(persona_key, 0) > 1:
                    errors.append("Médico sin usuario duplicado en el archivo (mismo nombre).")
                elif persona_key in existing_persona_keys:
                    nombre_completo = " ".join(
                        p for p in (parsed["paternoRaw"], parsed["maternoRaw"], parsed["nombreRaw"]) if p
                    )
                    errors.append(
                        f"Ya existe un médico sin usuario con el nombre '{nombre_completo}'."
                    )
            if not parsed["cedulaRaw"]:
                warnings.append("Sin Usuario y sin Cédula Profesional: el médico queda sin cédula.")
        else:
            usuario_obj = usuarios_by_username.get(parsed["usuario"])
            if usuario_obj is None:
                errors.append(
                    f"No existe un usuario con nombre de usuario '{parsed['usuario']}'."
                )
            elif usuario_seen.get(parsed["usuario"], 0) > 1:
                errors.append("Usuario duplicado en el archivo.")
            elif usuario_obj.id_usuario in existing_medico_usuario_ids:
                errors.append("Este usuario ya tiene un perfil de médico.")
            elif not MedicoRepository.usuario_es_medico(usuario_obj):
                # Misma regla que el alta manual (medico_usecase.create_medico).
                errors.append(
                    f"El usuario '{parsed['usuario']}' no tiene tipo de personal Médico."
                )
            if parsed["paternoRaw"] or parsed["maternoRaw"] or parsed["nombreRaw"]:
                warnings.append(
                    "Con Usuario, los apellidos y el nombre del Excel se ignoran "
                    "(se toman del usuario)."
                )

        # Cédula (identificador: no se trunca)
        if len(parsed["cedulaRaw"]) > _CEDULA_MAX_LEN:
            errors.append(
                f"Cédula Profesional excede el máximo de {_CEDULA_MAX_LEN} caracteres."
            )

        # Clave Clínica -> centro
        centro_id = None
        clave_value = None
        if parsed["claveClinicaRaw"]:
            matches = []
            if parsed["claveClinicaRaw"].isdigit():
                clave_value = int(parsed["claveClinicaRaw"])
                matches = centros_by_clave.get(clave_value, [])
            if len(matches) == 1:
                centro_id = matches[0]
            elif not matches:
                warnings.append(
                    f"Clave Clínica '{parsed['claveClinicaRaw']}' sin coincidencia: "
                    "el médico se crea sin centro."
                )
            else:
                warnings.append(
                    f"Clave Clínica '{parsed['claveClinicaRaw']}' es ambigua "
                    "(varios centros): el médico se crea sin centro."
                )

        # Especialidad
        especialidad_id = None
        if parsed["especialidadRaw"]:
            matches = especialidades_by_key.get(_nombre_key(parsed["especialidadRaw"]), [])
            if len(matches) == 1:
                especialidad_id = matches[0]
            elif not matches:
                warnings.append(
                    f"Especialidad '{parsed['especialidadRaw']}' sin coincidencia: "
                    "el médico se crea sin especialidad."
                )
            else:
                warnings.append(
                    f"Especialidad '{parsed['especialidadRaw']}' es ambigua: "
                    "el médico se crea sin especialidad."
                )

        tipo_medico_value, tipo_medico_valid = _resolve_choice(
            parsed["tipoMedicoRaw"], _TIPO_MEDICO_LOOKUP, DEFAULT_TIPO_MEDICO
        )
        if not tipo_medico_valid:
            errors.append(
                f"Tipo de Médico '{parsed['tipoMedicoRaw']}' no es válido. "
                f"Valores permitidos: {', '.join(code for code, _ in TIPO_MEDICO)}."
            )

        estatus_value, estatus_valid = _resolve_choice(
            parsed["estatusRaw"], _ESTATUS_MEDICO_LOOKUP, DEFAULT_ESTATUS_MEDICO
        )
        if not estatus_valid:
            errors.append(
                f"Estatus del Médico '{parsed['estatusRaw']}' no es válido. "
                f"Valores permitidos: {', '.join(code for code, _ in ESTATUS_MEDICO)} "
                "(o A / B)."
            )

        sin_usuario = not parsed["usuario"]
        row_data = {
            "noMedico": no_medico,
            "usuario": parsed["usuario"],
            "usuarioId": usuario_obj.id_usuario if usuario_obj else None,
            # Datos personales: solo se persisten para médicos sin usuario.
            "paterno": (_truncate(parsed["paternoRaw"], _NOMBRE_PARTE_MAX_LEN) or None) if sin_usuario else None,
            "materno": (_truncate(parsed["maternoRaw"], _NOMBRE_PARTE_MAX_LEN) or None) if sin_usuario else None,
            "nombre": (_truncate(parsed["nombreRaw"], _NOMBRE_PARTE_MAX_LEN) or None) if sin_usuario else None,
            # La cédula se conserva siempre: sin usuario va a CatMedico; con
            # usuario se crea en DetUsuarioCedula solo si no tiene ninguna.
            "cedulaProfesional": parsed["cedulaRaw"] or None,
            "claveClinica": clave_value,
            "centroId": centro_id,
            "especialidad": parsed["especialidadRaw"] or None,
            "especialidadId": especialidad_id,
            "tipoMedico": tipo_medico_value,
            "servicio": _truncate(parsed["servicioRaw"], _SERVICIO_MAX_LEN) or None,
            "estatusMedico": estatus_value,
            "observaciones": parsed["observacionesRaw"] or None,
        }

        result_rows.append(
            {
                "row": row_idx + 2,  # +1 header, +1 para 1-based
                "data": row_data,
                "errors": errors,
                "warnings": warnings,
            }
        )
        if errors:
            total_errores += 1
        if warnings:
            total_warnings += 1

    return {
        "total_records": total_records,
        "total_errores": total_errores,
        "total_warnings": total_warnings,
        "rows": result_rows,
    }


def _read_dataframe(file) -> pd.DataFrame:
    filename = (getattr(file, "name", "") or "").lower()
    if not filename.endswith((".xlsx", ".xls")):
        raise ImportFileError(
            code="IMPORT_FILE_INVALID",
            message="El archivo debe ser .xlsx o .xls.",
        )

    try:
        df = pd.read_excel(file, dtype=str, engine="openpyxl" if filename.endswith(".xlsx") else None)
    except Exception as exc:
        raise ImportFileError(
            code="IMPORT_FILE_INVALID",
            message="No se pudo leer el archivo. Verifique que sea un Excel válido.",
        ) from exc

    df = df.fillna("")

    actual_headers = [str(c).strip() for c in df.columns.tolist()]
    if actual_headers != HEADERS:
        raise ImportFileError(
            code="IMPORT_HEADERS_MISMATCH",
            message="Las columnas del archivo no coinciden con la plantilla de médicos.",
            details={"expected": HEADERS, "actual": actual_headers},
        )

    if len(df) > MAX_ROWS:
        raise ImportFileError(
            code="IMPORT_TOO_MANY_ROWS",
            message=f"El archivo excede el máximo de {MAX_ROWS} filas.",
        )

    return df
